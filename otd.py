"""准交率计算（收料单驱动口径）

口径（2026-10-19 用户确认）：
  1) 数据源：收料通知单（src/otd_recv.xlsx），按「源单单号 + 物料编码」回连全量采购订单，取交货日期/创建人/供应商
  2) 范围：仅统计 10 位采购订单创建人；供应商须在产能表（40 家箱包厂）内 → 自动排除包材/辅料
  3) 判定：收料日期 ≤ 采购订单交货日期 → 准时
  4) 准交率 = Σ准时收料数量 ÷ Σ收料总数量（分母是收料量，不是订单量）
  5) 数量：优先「合格入库关联数量」；为空且行状态=正常时取「实收数量」

输出 data.json.otd
"""
import calendar
import datetime as dt
import os
import re
from collections import defaultdict

import openpyxl

KNOWN_EXCLUDED = {
    "CGDD012692": "深圳市华熠网络·京东平台（固资）",
    "CGDD012720": "深圳市华熠网络·京东平台（固资）",
    "CGDD012843": "深圳市华熠网络·京东平台（固资）",
    "CGDD013390": "深圳市华熠网络·永隆达电脑（固资）",
}

# 仅保留这 10 位采购订单创建人（用户指定名单）
ALLOWED_CREATORS = [
    "吴巧玲", "张世豪", "聂婷", "梁成杰", "程美琪",
    "孔志豪", "姚冰婷", "张洁", "李淑云", "何君豪",
]
ALLOWED_SET = set(ALLOWED_CREATORS)
EXCLUDED_CREATORS = ["谭丁香", "曾琨", "曾坤", "谢国超"]

ALLOWED_STATUS = {"已审核", "重新审核"}
_STATUS_RANK = {"已审核": 3, "重新审核": 2, "审核中": 1, "创建": 0, "": 0}


def _load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    wb.close()
    hdr = [str(c).strip() if c is not None else "" for c in rows[0]]
    return hdr, rows[1:]


def _num(v):
    try:
        return float(str(v).replace(",", ""))
    except Exception:
        return 0.0


def _s(v):
    return "" if v is None else str(v).strip()


def _date(v):
    """兼容 2026/10/1、2026-10-01、datetime 等格式（金蝶导出月日不补零）"""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    t = _s(v).replace(" 00:00:00", "").strip()[:10]
    if not t:
        return None
    for sep in ("/", "-", "."):
        if sep in t:
            parts = t.split(sep)
            if len(parts) >= 3:
                try:
                    return dt.date(int(parts[0]), int(parts[1]), int(parts[2][:2]))
                except Exception:
                    return None
    return None


def _norm_creator(v):
    """张世豪01 -> 张世豪；去掉尾部数字与空格"""
    return re.sub(r"\d+$", "", _s(v)).strip()


def _find_src(src_dir, exact, keyword):
    p = os.path.join(src_dir, exact)
    if os.path.exists(p):
        return p
    cands = []
    if os.path.isdir(src_dir):
        for fn in os.listdir(src_dir):
            if fn.lower().endswith(".xlsx") and not fn.startswith("~$") and keyword in fn:
                fp = os.path.join(src_dir, fn)
                cands.append((os.path.getmtime(fp), fp))
    return max(cands)[1] if cands else None


def _month_bounds(period):
    y, m = int(period[:4]), int(period[5:7])
    return dt.date(y, m, 1), dt.date(y, m, calendar.monthrange(y, m)[1])


def build_otd(src_dir, periods=None):
    recv_path = _find_src(src_dir, "otd_recv.xlsx", "收料通知单")
    po_path = _find_src(src_dir, "otd_po_full.xlsx", "采购订单_全量")
    cap_path = _find_src(src_dir, "产能.xlsx", "产能")
    if not recv_path or not po_path:
        return None

    rh, rrows = _load(recv_path)
    ridx = {h: i for i, h in enumerate(rh)}

    def rg(r, name):
        i = ridx.get(name)
        return r[i] if i is not None and i < len(r) else None

    ph, prows = _load(po_path)
    pidx = {h: i for i, h in enumerate(ph)}

    def pg(r, name):
        i = pidx.get(name)
        return r[i] if i is not None and i < len(r) else None

    # ---------- 产能表：箱包成品供应商白名单（排除包材/辅料） ----------
    cap_sups = set()
    cap_short = {}   # 全称 -> 简称（前端核心/普通筛选要用）
    if cap_path:
        ch, crows = _load(cap_path)
        cidx = {h: i for i, h in enumerate(ch)}
        si = cidx.get("供应商简称")
        for r in crows:
            i = cidx.get("供应商")
            if i is not None and i < len(r):
                v = _s(r[i])
                if v:
                    cap_sups.add(v)
                    if si is not None and si < len(r):
                        cap_short[v] = _s(r[si]) or v

    # ---------- 采购订单：(单号,物料) -> 交货日期 / 创建人 / 供应商 ----------
    # 「单据类型」列可选：全量方案带该列时，准交率只统计「标准采购订单」，排除 VMI。
    has_doc_type = "单据类型" in pidx
    po = {}
    po_docs = set()
    po_by_doc = defaultdict(list)
    for r in prows:
        doc = _s(pg(r, "单据编号"))
        if not doc:
            continue
        po_docs.add(doc)
        mc = _s(pg(r, "物料编码"))
        st = _s(pg(r, "单据状态"))
        rank = _STATUS_RANK.get(st, 0)
        key = (doc, mc)
        rec = {
            "_rank": rank,
            "doc_type": _s(pg(r, "单据类型")),
            "delivery": _date(pg(r, "交货日期")),
            "creator": _norm_creator(pg(r, "创建人")),
            "creator_raw": _s(pg(r, "创建人")),
            "supplier": _s(pg(r, "供应商")),
            "status": st,
            "purchase_date": _date(pg(r, "采购日期")),
        }
        if key not in po or rank > po[key]["_rank"]:
            po[key] = rec
        po_by_doc[doc].append(rec)

    # ---------- 收料行 ----------
    receipts = []
    receipt_dates = []
    for r in rrows:
        dd = _date(rg(r, "收料日期"))
        if not dd:
            continue
        q_assoc = rg(r, "合格入库关联数量")
        q_actual = rg(r, "实收数量")
        if _s(q_assoc) not in ("", "None"):
            qty = _num(q_assoc)
        elif _s(rg(r, "行状态")) == "正常":
            qty = _num(q_actual)
        else:
            qty = 0.0
        if qty <= 0:
            continue
        _mat = _s(rg(r, "物料编码"))
        if "VYTK" in _mat.upper():   # 全局排除（2026-10-20 用户指定）
            continue
        receipts.append({
            "doc": _s(rg(r, "源单单号")),
            "recv_doc": _s(rg(r, "单据编号")),
            "material": _mat,
            "supplier_recv": _s(rg(r, "供应商")),
            "buyer_recv": _norm_creator(rg(r, "采购员")),
            "date": dd,
            "qty": qty,
        })
        receipt_dates.append(dd)

    if not receipts:
        return None

    # ---------- 统计周期（按收料日期所在月） ----------
    today = dt.date.today()
    months = sorted({r["date"].strftime("%Y-%m") for r in receipts})
    # 仅统计当月（口径：收料通知单方案已固定为「收料日期=本月」）
    current_period = today.strftime("%Y-%m")
    if periods is None:
        periods = [current_period]

    period_data = []
    for period in periods:
        start, end = _month_bounds(period)
        p_rows = [r for r in receipts if start <= r["date"] <= end]
        if not p_rows:
            continue

        # 覆盖率（源单能否在采购订单里找到）
        uniq_docs = {r["doc"] for r in p_rows if r["doc"]}
        matched_docs = {d for d in uniq_docs if d in po_docs}
        rows_matched = sum(1 for r in p_rows if r["doc"] in po_docs)

        details = []
        sup = defaultdict(lambda: {"recv_qty": 0.0, "ontime": 0.0, "rows": 0, "late_rows": 0})
        buy = defaultdict(lambda: {"recv_qty": 0.0, "ontime": 0.0, "rows": 0, "late_rows": 0})
        excluded = defaultdict(float)  # 未计入原因 -> 数量
        tot = 0.0
        okt = 0.0

        for r in p_rows:
            v = po.get((r["doc"], r["material"]))
            if not v:
                cands = po_by_doc.get(r["doc"]) or []
                v = cands[0] if cands else None
            if not v:
                excluded["源单未匹配采购订单"] += r["qty"]
                continue
            if v["creator"] not in ALLOWED_SET:
                excluded["非指定采购员"] += r["qty"]
                continue
            # 只统计标准采购订单；VMI 采购订单是另一种结算方式，不计入准交率
            if has_doc_type and v.get("doc_type") and v["doc_type"] != "标准采购订单":
                excluded["非标准采购订单（VMI 等）"] += r["qty"]
                continue
            if cap_sups and v["supplier"] not in cap_sups:
                excluded["包材/辅料（非产能表供应商）"] += r["qty"]
                continue

            delivery = v["delivery"]
            if delivery:
                ontime = r["date"] <= delivery
                days = (r["date"] - delivery).days
                verdict = "准时" if ontime else "逾期"
            else:
                ontime = False
                days = None
                verdict = "无交期"

            tot += r["qty"]
            if ontime:
                okt += r["qty"]

            key = v["supplier"] or "（未知）"
            b = sup[key]
            b["short"] = cap_short.get(key, key)
            b["recv_qty"] += r["qty"]
            b["rows"] += 1
            if ontime:
                b["ontime"] += r["qty"]
            else:
                b["late_rows"] += 1
            bk = v["creator"]
            b2 = buy[bk]
            b2["recv_qty"] += r["qty"]
            b2["rows"] += 1
            if ontime:
                b2["ontime"] += r["qty"]
            else:
                b2["late_rows"] += 1

            details.append({
                "doc": r["doc"],
                "recv_doc": r["recv_doc"],
                "material": r["material"],
                "supplier": v["supplier"],
                "creator": v["creator"],
                "delivery": delivery.isoformat() if delivery else "",
                "receipt_date": r["date"].isoformat(),
                "days": days,
                "days_txt": (("提前 %d 天" % (-days)) if days is not None and days < 0
                             else (("逾期 %d 天" % days) if days is not None else "无交期")),
                "qty": round(r["qty"]),
                "verdict": verdict,
            })

        def _pack(d):
            out = []
            for name, v in d.items():
                rate = (v["ontime"] / v["recv_qty"] * 100) if v["recv_qty"] > 0 else 0.0
                out.append({
                    "name": name,
                    "short": v.get("short") or name,
                    "recv_qty": round(v["recv_qty"]),
                    "ontime": round(v["ontime"]),
                    "late": round(v["recv_qty"] - v["ontime"]),
                    "rate": round(rate, 2),
                    "rows": v["rows"],
                    "late_rows": v["late_rows"],
                })
            out.sort(key=lambda x: -x["recv_qty"])
            return out

        rate = (okt / tot * 100) if tot > 0 else 0.0
        total_all = tot + sum(excluded.values())
        details.sort(key=lambda x: (x["verdict"] != "逾期", -x["qty"]))

        # 按收料日期逐日聚合（供右列趋势图；点击可下钻当日明细）
        by_day = defaultdict(lambda: {"recv": 0.0, "ontime": 0.0, "rows": 0})
        for x in details:
            b = by_day[x["receipt_date"]]
            b["recv"] += x["qty"]
            b["rows"] += 1
            if x["verdict"] == "准时":
                b["ontime"] += x["qty"]
        trend = [{
            "date": k,
            "recv": round(v["recv"]),
            "ontime": round(v["ontime"]),
            "rate": round(v["ontime"] / v["recv"] * 100, 2) if v["recv"] else 0.0,
            "rows": v["rows"],
        } for k, v in sorted(by_day.items())]

        # 缺失源单（数据质量校验）
        miss = []
        for doc in sorted(uniq_docs - matched_docs):
            miss.append({
                "doc": doc,
                "reason": KNOWN_EXCLUDED.get(doc, "采购订单未匹配"),
                "known": doc in KNOWN_EXCLUDED,
            })
        unknown_missing = [m for m in miss if not m["known"]]

        period_data.append({
            "period": period,
            "label": f"{period[:4]}年{int(period[5:7])}月",
            "valid": period == current_period,
            "validity_note": (
                "按收料日期统计；收料日期 ≤ 采购订单交货日期 记为准时。"
                if period == current_period
                else "按收料日期统计，历史月份数值完整可对比。"
            ),
            "overall": {
                "recv_qty": round(tot),
                "ontime": round(okt),
                "late": round(tot - okt),
                "rate": round(rate, 2),
                "receipt_rows": len(p_rows),
                "included_rows": len(details),
                "late_rows": sum(1 for x in details if x["verdict"] != "准时"),
                "excluded_qty": round(sum(excluded.values())),
                "excluded_rate": round(sum(excluded.values()) / total_all * 100, 2) if total_all else 0.0,
            },
            "coverage": {
                "unique_total": len(uniq_docs),
                "unique_matched": len(matched_docs),
                "unique_rate": round(len(matched_docs) / len(uniq_docs) * 100, 2) if uniq_docs else 0.0,
                "rows_total": len(p_rows),
                "rows_matched": rows_matched,
                "row_rate": round(rows_matched / len(p_rows) * 100, 2) if p_rows else 0.0,
                "receipt_window_start": min(receipt_dates).isoformat(),
            },
            "suppliers": _pack(sup),
            "buyers": _pack(buy),
            # 只保留当前月逐条明细（历史月仅保留汇总，避免单文件膨胀 >1MB）
            "order_details": details if period == current_period else [],
            "trend": trend if period == current_period else [],
            "excluded_reasons": [{"reason": k, "qty": round(v)} for k, v in
                                 sorted(excluded.items(), key=lambda x: -x[1])],
            "supplier_pool": len(cap_sups),
            "missing": miss,
            "unknown_missing": unknown_missing,
            "excluded_known": [m for m in miss if m["known"]],
            "blocked": False,
        })

    if not period_data:
        return None

    return {
        "updated_at": dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).strftime("%Y-%m-%d %H:%M"),
        "source_recv": os.path.basename(recv_path),
        "source_po_full": os.path.basename(po_path),
        "metric": "receipt_ontime_ratio",
        "metric_label": "收料准时率（收料日期≤交货日期）",
        "metric_formula": "Σ收料日期≤交货日期的收料量 ÷ Σ收料总量（仅标准采购订单，剔除 VMI）",
        "vmi_filter_enabled": has_doc_type,
        "creator_whitelist": ALLOWED_CREATORS,
        "excluded_creators": EXCLUDED_CREATORS,
        "supplier_pool": len(cap_sups),
        "receipt_window_start": min(receipt_dates).isoformat(),
        "default_period": (current_period if any(p["period"] == current_period for p in period_data)
                           else period_data[-1]["period"]),
        "periods": period_data,
        "known_excluded": KNOWN_EXCLUDED,
    }
