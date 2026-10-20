# -*- coding: utf-8 -*-
"""
构建管理台数据文件（console/weekly.json、console/apply.csv）。

口径（2026-10-20 起）：**申请单驱动**
  - 数据源：`src/otd_apply.xlsx`（金蝶「采购申请单列表 → 周计划申请(共享)」，每日同步）
  - 待下单数量 = 采购申请单「剩余数量」（= 批准数量 − 已下单数量）
  - 只取「剩余数量 > 0」的行
  - 供应商：① console/map.json（物料编码 → 供应商全称/简称/采购员）
            ② 采购订单历史（物料编码 → 供应商）兜底
            ③ 都没有 → 归入「待填写」分组，等人工指定
  - 明细保留：单据编号（CGSQ/JDYPP/JDYPR…）、领星创建人、规格型号、采购员、申请/批准数量

旧的 `src/周计划9.22.xlsx` 不再参与待下单计算（9/22 快照已过期）。
"""
import csv
import datetime as dt
import json
import os
import re

try:
    from clean import SHORT_OVERRIDE, BUYER_OVERRIDE, is_excluded_sku
except Exception:  # 独立运行时兜底
    SHORT_OVERRIDE, BUYER_OVERRIDE = {}, {}

    def is_excluded_sku(code):
        return "VYTK" in str(code or "").upper()

try:
    import openpyxl
except ImportError:
    openpyxl = None

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "src")
CONSOLE = os.path.join(HERE, "console")
PENDING = "待填写"


def norm_code(s):
    t = re.sub(r"[^A-Za-z0-9]", "", str(s or "")).upper()
    return re.sub(r"^SP(?=[A-Z])", "", t)


def raw_code(s):
    return re.sub(r"[^A-Za-z0-9]", "", str(s or "")).upper()


def num(v):
    try:
        return float(str(v).replace(",", "").strip())
    except Exception:
        return 0.0


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    t = str(v or "").strip()[:10].replace("/", "-")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})$", t)
    if m:
        return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def load_map_index():
    """map.json: [[物料编码, 供应商全称, 采购员, 供应商简称], ...]"""
    idx = {}
    p = os.path.join(CONSOLE, "map.json")
    if not os.path.exists(p):
        return idx
    try:
        rows = json.load(open(p, encoding="utf-8"))
    except Exception:
        return idx
    for row in rows:
        if not isinstance(row, (list, tuple)) or len(row) < 4:
            continue
        code = raw_code(row[0])
        if code and code not in idx:
            idx[code] = {"name": str(row[1] or "").strip(), "buyer": str(row[2] or "").strip(),
                         "short": str(row[3] or "").strip()}
    return idx


def load_po_supplier_index():
    """采购订单历史：物料编码 -> 供应商（map.json 未覆盖时兜底）"""
    idx = {}
    p = os.path.join(SRC, "采购订单.xlsx")
    if not os.path.exists(p) or openpyxl is None:
        return idx
    try:
        wb = openpyxl.load_workbook(p, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return idx
    h = [str(c).strip() if c is not None else "" for c in rows[0]]
    i = {k: n for n, k in enumerate(h)}
    if "物料编码" not in i or "供应商" not in i:
        return idx
    for r in rows[1:]:
        code = raw_code(r[i["物料编码"]])
        sup = str(r[i["供应商"]] or "").strip()
        if code and sup and code not in idx:
            idx[code] = {"name": sup, "short": "", "buyer": ""}
    return idx


def load_cap_short():
    """产能表：供应商全称 -> 简称"""
    out = {}
    p = os.path.join(SRC, "产能.xlsx")
    if not os.path.exists(p) or openpyxl is None:
        return out
    try:
        wb = openpyxl.load_workbook(p, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return out
    h = [str(c).strip() if c is not None else "" for c in rows[0]]
    i = {k: n for n, k in enumerate(h)}
    if "供应商" not in i:
        return out
    for r in rows[1:]:
        full = str(r[i["供应商"]] or "").strip()
        short = str(r[i.get("供应商简称", -1)] or "").strip() if "供应商简称" in i else ""
        if full:
            out[full] = short or full
    return out


def build_kind_index(map_idx):
    """款号(编码内数字前 4 位) -> 供应商，末行优先（与管理台 buildMapIndex 口径一致）"""
    kind = {}
    for code, info in map_idx.items():
        m = re.search(r"\d{4}", code)
        if m and info.get("name"):
            kind[m.group(0)] = info
    return kind


def resolve_supplier(code_raw, map_idx, kind_idx):
    """返回 (info, via)：exact / pfx / kind / ''（与管理台 lookupSKU 同序）"""
    code = raw_code(code_raw)
    if not code:
        return None, ""
    hit = map_idx.get(code)
    if hit:
        return hit, "exact"
    for pfx in ("SP",):
        if code.startswith(pfx) and code[len(pfx):] in map_idx:
            return map_idx[code[len(pfx):]], "pfx"
    m = re.search(r"\d{4}", code)
    if m and m.group(0) in kind_idx:
        return kind_idx[m.group(0)], "kind"
    return None, ""


def build():
    os.makedirs(CONSOLE, exist_ok=True)
    core = set()
    updated_at = None
    try:
        data = json.load(open(os.path.join(HERE, "data.json"), encoding="utf-8"))
        core = set((data.get("meta") or {}).get("core_suppliers") or [])
        updated_at = (data.get("meta") or {}).get("updated_at")
    except Exception:
        pass

    apply_path = os.path.join(SRC, "otd_apply.xlsx")
    if not os.path.exists(apply_path) or openpyxl is None:
        print("⚠️ 缺少 src/otd_apply.xlsx 或 openpyxl，跳过待下单数据")
        return

    map_idx = load_map_index()
    kind_idx = build_kind_index(map_idx)
    po_idx = load_po_supplier_index()
    cap_short = load_cap_short()
    via_count = {}
    pending_skus = {}

    wb = openpyxl.load_workbook(apply_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    h = [str(c).strip() if c is not None else "" for c in rows[0]]
    i = {k: n for n, k in enumerate(h)}

    def g(r, *keys):
        for k in keys:
            n = i.get(k)
            if n is not None and n < len(r):
                v = r[n]
                if v is not None and str(v).strip() != "":
                    return v
        return None

    groups = {}
    hit_map = hit_po = miss = 0
    total_qty = 0.0
    skipped_zero = 0
    skipped_bad = 0
    skipped_vytk = 0

    for r in rows[1:]:
        remain = num(g(r, "剩余数量"))
        if remain <= 0:
            skipped_zero += 1
            continue
        code_raw = str(g(r, "物料编码") or "").strip()
        # 跳过源表末尾的「合计」行（物料编码为空、单据编号=合计）
        if not code_raw:
            skipped_bad += 1
            continue
        if is_excluded_sku(code_raw):
            skipped_vytk += 1
            continue
        code = raw_code(code_raw)
        info, via = resolve_supplier(code_raw, map_idx, kind_idx)
        src = via or ""
        if not info:
            info = po_idx.get(code)
            src = "po" if info else ""
        if info:
            full = info.get("name") or ""
            short = SHORT_OVERRIDE.get(full) or info.get("short") or cap_short.get(full) or full
            buyer = BUYER_OVERRIDE.get(full) or info.get("buyer") or str(g(r, "采购员") or "").strip()
            if src in ("exact", "pfx", "kind"):
                hit_map += 1
                via_count[src] = via_count.get(src, 0) + 1
            else:
                hit_po += 1
        else:
            full, short, buyer = "", PENDING, str(g(r, "采购员") or "").strip()
            miss += 1
            pending_skus[code_raw] = pending_skus.get(code_raw, 0) + remain

        key = short or full or PENDING
        grp = groups.setdefault(key, {
            "name": full or ("未匹配供应商（待人工填写）" if key == PENDING else key),
            "short": key,
            "core": key in core,
            "pending": key == PENDING,
            "plan": 0.0,
            "skus": set(),
            "spus": set(),
            "buyers": set(),
            "docs": set(),
            "items": [],
        })
        qty = round(remain, 1)
        total_qty += qty
        grp["plan"] += qty
        if code_raw:
            grp["skus"].add(code_raw)
        spu = "".join(ch for ch in code_raw if ch.isdigit())[:4]
        if spu:
            grp["spus"].add(spu)
        if buyer:
            grp["buyers"].add(buyer)
        doc = str(g(r, "单据编号") or "").strip()
        if doc:
            grp["docs"].add(doc)
        grp["items"].append({
            "sku": code_raw,
            "spu": spu,
            "name": str(g(r, "物料名称") or "").strip(),
            "spec": str(g(r, "规格型号") or "").strip(),
            "buyer": buyer,
            "qty": qty,
            "doc": doc,
            "ling": str(g(r, "领星创建人") or "").strip(),
            "date": (to_date(g(r, "创建日期")) or "").isoformat() if to_date(g(r, "创建日期")) else "",
            "apply_qty": round(num(g(r, "申请数量")), 1),
            "approved_qty": round(num(g(r, "批准数量")), 1),
            "ordered_qty": round(num(g(r, "订单数量")), 1),
            "status": str(g(r, "数据状态") or "").strip(),
            "mapped": bool(info),
        })

    suppliers = []
    for key, grp in groups.items():
        suppliers.append({
            "name": grp["name"],
            "short": grp["short"],
            "core": grp["core"],
            "pending": grp["pending"],
            "plan": round(grp["plan"], 1),
            "skus": len(grp["skus"]),
            "spus": len(grp["spus"]),
            "docs": len(grp["docs"]),
            "buyers": sorted(grp["buyers"]),
            "items": grp["items"],
        })
    suppliers.sort(key=lambda x: (not x["pending"], -x["plan"]))

    out = {
        "updated_at": updated_at,
        "source": "apply",
        "source_file": os.path.basename(apply_path),
        "total_plan": round(total_qty, 1),
        "rows": sum(len(x["items"]) for x in suppliers),
        "supplier_count": len([x for x in suppliers if not x["pending"]]),
        "core_count": len([x for x in suppliers if x["core"]]),
        "pending": {
            "rows": len([it for x in suppliers if x["pending"] for it in x["items"]]),
            "qty": round(sum(x["plan"] for x in suppliers if x["pending"]), 1),
        },
        "pending_skus": [{"sku": k, "qty": round(v, 1)} for k, v in
                         sorted(pending_skus.items(), key=lambda x: -x[1])],
        "apply_match": {"by_map": hit_map, "by_po": hit_po, "pending": miss, "via": via_count,
                        "skipped_zero": skipped_zero, "skipped_bad": skipped_bad,
                        "skipped_vytk": skipped_vytk},
        "suppliers": suppliers,
    }
    with open(os.path.join(CONSOLE, "weekly.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))

    # apply.csv 供管理台自动载入解析（金蝶 xlsx 在内联 SheetJS 里会内存溢出）
    def cv(v):
        if isinstance(v, (dt.datetime, dt.date)):
            return "%d/%d/%d" % (v.year, v.month, v.day)
        return "" if v is None else v

    with open(os.path.join(CONSOLE, "apply.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        for r in rows:
            w.writerow([cv(x) for x in r])

    print("待下单（申请单驱动）：%d 行 / %s 件 ｜ 供应商 %d 家 + 待填写 %s 件（%d 行）｜ 映射 map %d / PO %d / 未映射 %d"
          % (out["rows"], out["total_plan"], out["supplier_count"], out["pending"]["qty"],
             out["pending"]["rows"], hit_map, hit_po, miss))


if __name__ == "__main__":
    build()
